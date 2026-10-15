import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from generate_test_plan import test_items

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Verification Test Plan"

# Gridlines
ws.views.sheetView[0].showGridLines = True

# Colors
COLOR_SUMMARY_BG = "FFF2CC" # Soft yellow
COLOR_GOLD       = "FFE699" # Gold/Yellow
COLOR_ORANGE     = "F8CBAD" # Soft Orange / Peach
COLOR_PEACH      = "FCE4D6" # Test Pattern Phase
COLOR_BLUE       = "D9E1F2" # Soft Blue
COLOR_GRAY_LIGHT = "F2F2F2"
COLOR_PASS_BG    = "E2EFDA" # Light Green
COLOR_PASS_TXT   = "375623" # Dark Green
COLOR_BORDER     = "D9D9D9" # Subtle grid or black
COLOR_BORDER_HDR = "000000"

FILL_SUMMARY     = PatternFill(start_color=COLOR_SUMMARY_BG, end_color=COLOR_SUMMARY_BG, fill_type="solid")
FILL_GOLD        = PatternFill(start_color=COLOR_GOLD, end_color=COLOR_GOLD, fill_type="solid")
FILL_ORANGE      = PatternFill(start_color=COLOR_ORANGE, end_color=COLOR_ORANGE, fill_type="solid")
FILL_PEACH       = PatternFill(start_color=COLOR_PEACH, end_color=COLOR_PEACH, fill_type="solid")
FILL_BLUE        = PatternFill(start_color=COLOR_BLUE, end_color=COLOR_BLUE, fill_type="solid")
FILL_GRAY        = PatternFill(start_color=COLOR_GRAY_LIGHT, end_color=COLOR_GRAY_LIGHT, fill_type="solid")
FILL_PASS        = PatternFill(start_color=COLOR_PASS_BG, end_color=COLOR_PASS_BG, fill_type="solid")

FONT_TITLE       = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_HEADER      = Font(name="Calibri", size=10, bold=True, color="000000")
FONT_SUBHEADER   = Font(name="Calibri", size=9.5, bold=True, color="000000")
FONT_BODY        = Font(name="Calibri", size=9, color="000000")
FONT_BODY_BOLD   = Font(name="Calibri", size=9, bold=True, color="000000")
FONT_CODE        = Font(name="Consolas", size=8.5, color="1F497D")
FONT_PASS        = Font(name="Calibri", size=9.5, bold=True, color=COLOR_PASS_TXT)
FONT_CRITICAL    = Font(name="Calibri", size=9, bold=True, color="C00000")
FONT_HIGH        = Font(name="Calibri", size=9, bold=True, color="ED7D31")
FONT_MEDIUM      = Font(name="Calibri", size=9, bold=True, color="7030A0")
FONT_DATE        = Font(name="Calibri", size=8.5, color="595959")

THIN_BLACK = Side(border_style="thin", color="000000")
THIN_GRAY  = Side(border_style="thin", color="BFBFBF")
DOTTED     = Side(border_style="dotted", color="A6A6A6")

BORDER_HEADER = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
BORDER_DATA   = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)

ALIGN_CENTER     = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT       = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_TOP_LEFT   = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_TOP_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

# -----------------------------------------------------------------
# 1. Summary Header (Row 1 & 2)
# -----------------------------------------------------------------
ws.merge_cells("A1:E1")
ws["A1"] = "Summary: Verification Progress"
ws["A1"].font = FONT_TITLE
ws["A1"].fill = FILL_SUMMARY
ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

total_items = len(test_items)
end_row = 5 + total_items

ws["F1"] = "Total Items:"
ws["F1"].font = FONT_HEADER
ws["F1"].alignment = ALIGN_CENTER
ws["F1"].fill = FILL_GRAY

ws["G1"] = f"=COUNTA(E6:E{end_row})"
ws["G1"].font = FONT_BODY_BOLD
ws["G1"].alignment = ALIGN_CENTER
ws["G1"].fill = FILL_ORANGE

ws["H1"] = "Finished:"
ws["H1"].font = FONT_HEADER
ws["H1"].alignment = ALIGN_CENTER
ws["H1"].fill = FILL_GRAY

ws["I1"] = f'=COUNTIF(N6:N{end_row}, "Pass")'
ws["I1"].font = FONT_BODY_BOLD
ws["I1"].alignment = ALIGN_CENTER
ws["I1"].fill = FILL_PASS

ws["J1"] = "Progress Rate:"
ws["J1"].font = FONT_HEADER
ws["J1"].alignment = ALIGN_CENTER
ws["J1"].fill = FILL_GRAY

ws["K1"] = "=I1/G1"
ws["K1"].number_format = "0.0%"
ws["K1"].font = FONT_BODY_BOLD
ws["K1"].alignment = ALIGN_CENTER
ws["K1"].fill = FILL_PASS

ws.merge_cells("A2:B2")
ws["A2"] = "Module name"
ws["A2"].font = FONT_HEADER
ws["A2"].fill = FILL_GOLD
ws["A2"].alignment = ALIGN_CENTER

ws.merge_cells("C2:E2")
ws["C2"] = "hms_timer (Top-Level Universal Architecture)"
ws["C2"].font = FONT_BODY_BOLD
ws["C2"].fill = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
ws["C2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

for r in [1, 2]:
    for c in range(1, 20):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER_HEADER
        if r == 1 and c > 11:
            cell.fill = FILL_GRAY
        elif r == 2 and c > 5:
            cell.fill = FILL_GRAY

# -----------------------------------------------------------------
# 2. Main Headers (Rows 3, 4, 5)
# -----------------------------------------------------------------
# Row 3 Super Headers
ws.merge_cells("A3:A5")
ws["A3"] = "Big item"
ws["A3"].font = FONT_HEADER
ws["A3"].fill = FILL_GOLD
ws["A3"].alignment = ALIGN_CENTER

ws.merge_cells("B3:B5")
ws["B3"] = "No"
ws["B3"].font = FONT_HEADER
ws["B3"].fill = FILL_GOLD
ws["B3"].alignment = ALIGN_CENTER

ws.merge_cells("C3:C5")
ws["C3"] = "Medium item"
ws["C3"].font = FONT_HEADER
ws["C3"].fill = FILL_GOLD
ws["C3"].alignment = ALIGN_CENTER

ws.merge_cells("D3:D5")
ws["D3"] = "No"
ws["D3"].font = FONT_HEADER
ws["D3"].fill = FILL_GOLD
ws["D3"].alignment = ALIGN_CENTER

ws.merge_cells("E3:E5")
ws["E3"] = "Small item"
ws["E3"].font = FONT_HEADER
ws["E3"].fill = FILL_GOLD
ws["E3"].alignment = ALIGN_CENTER

# Planning Phase: F3 to K3
ws.merge_cells("F3:K3")
ws["F3"] = "Planning Phase"
ws["F3"].font = FONT_HEADER
ws["F3"].fill = FILL_GOLD
ws["F3"].alignment = ALIGN_CENTER

# Test Pattern/SVA Implementation Phase: L3 to M3
ws.merge_cells("L3:M3")
ws["L3"] = "Test Pattern/ SVA Implementation Phase"
ws["L3"].font = FONT_HEADER
ws["L3"].fill = FILL_PEACH
ws["L3"].alignment = ALIGN_CENTER

# Execution Phase: N3 to S3
ws.merge_cells("N3:S3")
ws["N3"] = "Execution Phase"
ws["N3"].font = FONT_HEADER
ws["N3"].fill = FILL_BLUE
ws["N3"].alignment = ALIGN_CENTER

# Row 4 Sub-headers
ws.merge_cells("F4:F5")
ws["F4"] = "Description"
ws["F4"].font = FONT_HEADER
ws["F4"].fill = FILL_GOLD
ws["F4"].alignment = ALIGN_CENTER

ws.merge_cells("G4:G5")
ws["G4"] = "Expected"
ws["G4"].font = FONT_HEADER
ws["G4"].fill = FILL_GOLD
ws["G4"].alignment = ALIGN_CENTER

ws["H4"] = "Number"
ws["H4"].font = FONT_SUBHEADER
ws["H4"].fill = FILL_ORANGE
ws["H4"].alignment = ALIGN_CENTER

ws["I4"] = "Finished"
ws["I4"].font = FONT_SUBHEADER
ws["I4"].fill = FILL_ORANGE
ws["I4"].alignment = ALIGN_CENTER

# Row 5 Formula / Count Cells for Number & Finished
ws["H5"] = f"=COUNTA(E6:E{end_row})"
ws["H5"].font = FONT_BODY_BOLD
ws["H5"].fill = FILL_ORANGE
ws["H5"].alignment = ALIGN_CENTER

ws["I5"] = f'=COUNTIF(N6:N{end_row}, "Pass")'
ws["I5"].font = FONT_BODY_BOLD
ws["I5"].fill = FILL_ORANGE
ws["I5"].alignment = ALIGN_CENTER

ws.merge_cells("J4:J5")
ws["J4"] = "Priority"
ws["J4"].font = FONT_HEADER
ws["J4"].fill = FILL_GOLD
ws["J4"].alignment = ALIGN_CENTER

ws.merge_cells("K4:K5")
ws["K4"] = "PIC"
ws["K4"].font = FONT_HEADER
ws["K4"].fill = FILL_GOLD
ws["K4"].alignment = ALIGN_CENTER

ws.merge_cells("L4:L5")
ws["L4"] = "Test/Checker File"
ws["L4"].font = FONT_HEADER
ws["L4"].fill = FILL_PEACH
ws["L4"].alignment = ALIGN_CENTER

ws.merge_cells("M4:M5")
ws["M4"] = "Task Number/ SVA"
ws["M4"].font = FONT_HEADER
ws["M4"].fill = FILL_PEACH
ws["M4"].alignment = ALIGN_CENTER

# Execution Phase Rounds
ws.merge_cells("N4:P4")
ws["N4"] = "Round 1"
ws["N4"].font = FONT_HEADER
ws["N4"].fill = FILL_BLUE
ws["N4"].alignment = ALIGN_CENTER

ws["N5"] = "Result"
ws["N5"].font = FONT_SUBHEADER
ws["N5"].fill = FILL_BLUE
ws["N5"].alignment = ALIGN_CENTER

ws["O5"] = "Date"
ws["O5"].font = FONT_SUBHEADER
ws["O5"].fill = FILL_BLUE
ws["O5"].alignment = ALIGN_CENTER

ws["P5"] = "Note"
ws["P5"].font = FONT_SUBHEADER
ws["P5"].fill = FILL_BLUE
ws["P5"].alignment = ALIGN_CENTER

ws.merge_cells("Q4:S4")
ws["Q4"] = "Round 2"
ws["Q4"].font = FONT_HEADER
ws["Q4"].fill = FILL_BLUE
ws["Q4"].alignment = ALIGN_CENTER

ws["Q5"] = "Result"
ws["Q5"].font = FONT_SUBHEADER
ws["Q5"].fill = FILL_BLUE
ws["Q5"].alignment = ALIGN_CENTER

ws["R5"] = "Date"
ws["R5"].font = FONT_SUBHEADER
ws["R5"].fill = FILL_BLUE
ws["R5"].alignment = ALIGN_CENTER

ws["S5"] = "Note"
ws["S5"].font = FONT_SUBHEADER
ws["S5"].fill = FILL_BLUE
ws["S5"].alignment = ALIGN_CENTER

# Apply border to all header cells
for r in range(3, 6):
    for c in range(1, 20):
        ws.cell(row=r, column=c).border = BORDER_HEADER

# -----------------------------------------------------------------
# 3. Populate Data Rows
# -----------------------------------------------------------------
current_row = 6

for idx, item in enumerate(test_items):
    (b_no, b_item, m_no, m_item, s_no, s_item,
     muc_dich, stimulus, expected, priority, pic,
     test_file, task_sva, result, date_val, note_val) = item

    r = current_row

    # Col A: Big item
    ws.cell(row=r, column=1, value=b_item).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=1).font = FONT_BODY_BOLD

    # Col B: Big item No
    ws.cell(row=r, column=2, value=b_no).alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=2).font = FONT_BODY_BOLD

    # Col C: Medium item
    ws.cell(row=r, column=3, value=m_item).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=3).font = FONT_BODY_BOLD

    # Col D: Medium item No
    ws.cell(row=r, column=4, value=m_no).alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=4).font = FONT_BODY

    # Col E: Small item
    # Note: in the example image, Small item has No in a small cell or combined,
    # let's put Small item with its number and title
    ws.cell(row=r, column=5, value=f"{s_no}\n{s_item}").alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=5).font = FONT_BODY_BOLD

    # Col F: Description (Mục đích + Stimulus)
    desc_text = f"Mục đích: {muc_dich}\nStimulus: {stimulus}"
    ws.cell(row=r, column=6, value=desc_text).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=6).font = FONT_BODY

    # Col G: Expected
    ws.cell(row=r, column=7, value=expected).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=7).font = FONT_BODY

    # Col H: Number (count per small item is 1)
    ws.cell(row=r, column=8, value=1).alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=8).font = FONT_BODY

    # Col I: Finished (1 if pass)
    ws.cell(row=r, column=9, value=1 if result == "Pass" else 0).alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=9).font = FONT_BODY

    # Col J: Priority
    ws.cell(row=r, column=10, value=priority).alignment = ALIGN_TOP_CENTER
    if priority == "Critical":
        ws.cell(row=r, column=10).font = FONT_CRITICAL
    elif priority == "High":
        ws.cell(row=r, column=10).font = FONT_HIGH
    else:
        ws.cell(row=r, column=10).font = FONT_MEDIUM

    # Col K: PIC
    ws.cell(row=r, column=11, value=pic).alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=11).font = FONT_BODY

    # Col L: Test/Checker File
    ws.cell(row=r, column=12, value=test_file).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=12).font = FONT_CODE

    # Col M: Task Number/SVA
    ws.cell(row=r, column=13, value=task_sva).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=13).font = FONT_CODE

    # Col N: Round 1 Result
    c_res = ws.cell(row=r, column=14, value=result)
    c_res.alignment = ALIGN_TOP_CENTER
    if result == "Pass":
        c_res.fill = FILL_PASS
        c_res.font = FONT_PASS

    # Col O: Round 1 Date
    ws.cell(row=r, column=15, value=date_val).alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=15).font = FONT_DATE

    # Col P: Round 1 Note
    ws.cell(row=r, column=16, value=note_val).alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=16).font = FONT_BODY

    # Col Q: Round 2 Result
    ws.cell(row=r, column=17, value="").alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=17).font = FONT_BODY

    # Col R: Round 2 Date
    ws.cell(row=r, column=18, value="").alignment = ALIGN_TOP_CENTER
    ws.cell(row=r, column=18).font = FONT_DATE

    # Col S: Round 2 Note
    ws.cell(row=r, column=19, value="").alignment = ALIGN_TOP_LEFT
    ws.cell(row=r, column=19).font = FONT_BODY

    # Set borders and subtle background for alternate rows or clean white
    for c in range(1, 20):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER_DATA

    current_row += 1

# -----------------------------------------------------------------
# 4. Vertical Merging for Big Items and Medium Items
# -----------------------------------------------------------------
# Scan and merge Big item (Col 1 & 2)
big_start = 6
current_big = test_items[0][1]

for idx in range(1, len(test_items)):
    row_idx = 6 + idx
    item_big = test_items[idx][1]
    if item_big != current_big:
        # Merge previous big item
        if row_idx - 1 > big_start:
            ws.merge_cells(start_row=big_start, start_column=1, end_row=row_idx - 1, end_column=1)
            ws.merge_cells(start_row=big_start, start_column=2, end_row=row_idx - 1, end_column=2)
        big_start = row_idx
        current_big = item_big

# Final big item merge
if end_row > big_start:
    ws.merge_cells(start_row=big_start, start_column=1, end_row=end_row, end_column=1)
    ws.merge_cells(start_row=big_start, start_column=2, end_row=end_row, end_column=2)

# Scan and merge Medium item (Col 3 & 4)
med_start = 6
current_med = (test_items[0][1], test_items[0][3])

for idx in range(1, len(test_items)):
    row_idx = 6 + idx
    item_med = (test_items[idx][1], test_items[idx][3])
    if item_med != current_med:
        if row_idx - 1 > med_start:
            ws.merge_cells(start_row=med_start, start_column=3, end_row=row_idx - 1, end_column=3)
            ws.merge_cells(start_row=med_start, start_column=4, end_row=row_idx - 1, end_column=4)
        med_start = row_idx
        current_med = item_med

if end_row > med_start:
    ws.merge_cells(start_row=med_start, start_column=3, end_row=end_row, end_column=3)
    ws.merge_cells(start_row=med_start, start_column=4, end_row=end_row, end_column=4)

# -----------------------------------------------------------------
# 5. Set Optimal Column Widths and Row Heights
# -----------------------------------------------------------------
col_widths = {
    "A": 22, # Big item
    "B": 5,  # No
    "C": 26, # Medium item
    "D": 5,  # No
    "E": 28, # Small item
    "F": 46, # Description
    "G": 46, # Expected
    "H": 8,  # Number
    "I": 8,  # Finished
    "J": 10, # Priority
    "K": 13, # PIC
    "L": 22, # Test File
    "M": 26, # Task / SVA
    "N": 8,  # Round 1 Result
    "O": 11, # Round 1 Date
    "P": 32, # Round 1 Note
    "Q": 8,  # Round 2 Result
    "R": 11, # Round 2 Date
    "S": 20, # Round 2 Note
}

for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# Set row heights
ws.row_dimensions[1].height = 26
ws.row_dimensions[2].height = 24
ws.row_dimensions[3].height = 24
ws.row_dimensions[4].height = 22
ws.row_dimensions[5].height = 20

for r in range(6, end_row + 1):
    ws.row_dimensions[r].height = 68 # Enough space for multi-line description/expected

# Save workbook to docs/HMS_Timer_Test_Plan.xlsx
output_path = "docs/HMS_Timer_Test_Plan.xlsx"
wb.save(output_path)
print(f"Successfully generated {output_path} with {total_items} verification items.")
